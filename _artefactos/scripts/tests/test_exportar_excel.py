# -*- coding: utf-8 -*-
"""Tests de `exportar_excel.py` — export reproducible (CA-R2-10) · R2."""
from pathlib import Path
import csv
import sys

SCRIPTS = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(SCRIPTS))

import exportar_excel as E  # noqa: E402
import tabla_comun as TC  # noqa: E402


def _escribir(path, header, rows):
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", encoding="utf-8", newline="") as fh:
        w = csv.writer(fh, lineterminator="\n")
        w.writerow(header)
        w.writerows(rows)


def test_export_reproducible(tmp_path):
    hojas = tmp_path / "Hojas"
    _escribir(hojas / "T2_ataques.csv", ["n", "ata_id", "config", "umbral", "Detection"],
              [(1, "ATA001", "C1", 0, "SÍ")])
    # una fila de datos que EMPIEZA por '#' (T4) no debe tratarse como comentario
    _escribir(hojas / "T4_rendimiento.csv", ["metrica", "wazuh_C1_N0"],
              [("#Casos", 55), ("TPR", "0.963636")])
    a = tmp_path / "a.xlsx"
    b = tmp_path / "b.xlsx"
    E.exportar(str(hojas), str(a))
    E.exportar(str(hojas), str(b))
    assert a.read_bytes() == b.read_bytes()
    # y la fila '#Casos' está presente en el xlsx
    from openpyxl import load_workbook
    wb = load_workbook(str(a))
    ws = wb["T4_rendimiento"]
    vals = [ws.cell(row=r, column=1).value for r in range(1, ws.max_row + 1)]
    assert "#Casos" in vals
